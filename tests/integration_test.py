import asyncio
from hashlib import sha256
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import time
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document

from case_indexer.core import CaseIndexer, sha256_file
from server import get_index_status, read_transcript, start_background_index


class CaseIndexIntegrationTests(unittest.TestCase):
    def setUp(self) -> None:
        work = ROOT / "test-run"
        work.mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=work)
        self.project = Path(self.temporary.name) / "synthetic-project"
        self.project.mkdir()
        self.text = self.project / "alpha.txt"
        self.text.write_text("Synthetic alpha document.\nNo real case data.\n", encoding="utf-8")
        document = Document()
        document.add_paragraph("Synthetic beta document.")
        document.add_paragraph("Generated solely for integration testing.")
        self.docx = self.project / "beta.docx"
        document.save(self.docx)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def source_hashes(self) -> dict[str, str]:
        return {path.name: sha256_file(path) for path in (self.text, self.docx) if path.exists()}

    def test_incremental_lifecycle_and_source_preservation(self) -> None:
        originals = self.source_hashes()
        indexer = CaseIndexer(str(self.project))
        first = indexer.index(use_fallback=False)
        self.assertTrue(first["ok"], first)
        self.assertEqual(first["processed"], 2)
        self.assertEqual(first["documents_total"], 2)

        second = indexer.index(use_fallback=False)
        self.assertEqual(second["skipped"], 2)
        self.assertEqual(second["processed"], 0)

        self.text.write_text("Synthetic alpha document, changed.\n", encoding="utf-8")
        third = indexer.index(use_fallback=False)
        self.assertEqual(third["processed"], 1)
        self.assertEqual(third["skipped"], 1)

        transcript = read_transcript(str(self.project), "alpha.txt")
        self.assertIn("changed", transcript["text"])
        page = read_transcript(str(self.project), "alpha.txt", page=1)
        self.assertIn("changed", page["text"])
        status = get_index_status(str(self.project))
        self.assertEqual(len(status["documents"]), 2)
        self.assertIsNone(status["active_run"])

        self.docx.unlink()
        fourth = indexer.index(use_fallback=False)
        self.assertEqual(fourth["removed"], 1)
        self.assertEqual(fourth["documents_total"], 1)
        expected = {self.text.name: sha256_file(self.text)}
        self.assertEqual(self.source_hashes(), expected)
        self.assertTrue(originals)

    def test_lock_is_reported_without_breaking_index(self) -> None:
        indexer = CaseIndexer(str(self.project))
        indexer.index_dir.mkdir(parents=True)
        with indexer.lock:
            result = CaseIndexer(str(self.project)).index(use_fallback=False)
        self.assertFalse(result["ok"])
        self.assertEqual(result["error"], "already_running")

    def test_background_run_and_status(self) -> None:
        started = start_background_index(str(self.project), False, False)
        self.assertTrue(started["ok"], started)
        deadline = time.monotonic() + 60
        while time.monotonic() < deadline:
            status = get_index_status(str(self.project))
            if status["run"] and status["run"].get("status") == "completed":
                break
            time.sleep(0.1)
        else:
            self.fail("Background indexing did not finish within 60 seconds")
        self.assertEqual(status["run"]["counts"]["processed"], 2)
        self.assertTrue((self.project / "_work" / "case-index" / "manifest.json").is_file())
        # The status file is committed just before the worker closes its log and exits.
        while CaseIndexer._pid_alive(started["pid"]) and time.monotonic() < deadline:
            time.sleep(0.1)
        self.assertFalse(CaseIndexer._pid_alive(started["pid"]), "Background worker is still running")


class MCPStdioTests(unittest.TestCase):
    def test_all_tools_are_exposed_over_stdio(self) -> None:
        async def check() -> None:
            from mcp import ClientSession, StdioServerParameters
            from mcp.client.stdio import stdio_client

            parameters = StdioServerParameters(command=sys.executable, args=[str(ROOT / "server.py")])
            async with stdio_client(parameters) as (reader, writer):
                async with ClientSession(reader, writer) as session:
                    await session.initialize()
                    response = await session.list_tools()
                    names = {tool.name for tool in response.tools}
                    self.assertEqual(names, {"index_project", "cancel_index", "get_index_status", "read_transcript"})

        asyncio.run(asyncio.wait_for(check(), timeout=30))


if __name__ == "__main__":
    unittest.main(verbosity=2)
